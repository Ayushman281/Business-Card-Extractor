from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from app.models.lead import FIELDS, Lead

HEADERS = ("First Name", "Last Name", "Position / Job Title", "Company", "Location", "Phone Number", "Email Address")


def generate_excel(leads: list[Lead]) -> bytes:
    rows = [[getattr(lead, field) or "" for field in FIELDS] for lead in leads]
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(rows, columns=HEADERS).to_excel(writer, sheet_name="Leads", index=False)
        sheet = writer.sheets["Leads"]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="183D36")
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 28
        for column, width in zip("ABCDEFG", (20, 20, 30, 30, 44, 28, 36)):
            sheet.column_dimensions[column].width = width
        for row_index, values in enumerate(rows, start=2):
            for column_index, value in enumerate(values, start=1):
                cell = sheet.cell(row_index, column_index)
                # Explicit string type prevents Excel formulas, including values
                # beginning with =, +, -, or @. Preserve phone leading zeros.
                cell.value = value
                cell.data_type = "s"
                cell.number_format = "@"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    return output.getvalue()
