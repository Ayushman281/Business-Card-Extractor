from io import BytesIO

import pytest
from openpyxl import load_workbook
from PIL import Image

from app.core.config import Settings
from app.models.lead import Lead
from app.services.excel_service import HEADERS, generate_excel
from app.services.image_service import ImageValidationError, prepare_image, safe_filename


def test_jpeg_is_rgb_and_bounded(image_bytes):
    result = prepare_image(image_bytes, "card.JPG", Settings(_env_file=None))
    assert result.mode == "RGB"
    assert max(result.size) <= 1800
    result.close()


@pytest.mark.parametrize("name,data", [("card.gif", b"GIF89a"), ("card.png", b"corrupt"), ("card.jpg", b"")])
def test_bad_images(name, data):
    with pytest.raises(ImageValidationError):
        prepare_image(data, name, Settings(_env_file=None))


def test_extension_must_match_contents(image_bytes):
    with pytest.raises(ImageValidationError):
        prepare_image(image_bytes, "card.png", Settings(_env_file=None))


def test_pixel_limit_precedes_decode(image_bytes):
    with pytest.raises(ImageValidationError):
        prepare_image(image_bytes, "card.jpg", Settings(_env_file=None, max_image_pixels=10))


def test_exif_orientation():
    image = Image.new("RGB", (300, 100), "white")
    exif = Image.Exif()
    exif[274] = 6
    buffer = BytesIO()
    image.save(buffer, format="JPEG", exif=exif)
    corrected = prepare_image(buffer.getvalue(), "card.jpg", Settings(_env_file=None))
    assert corrected.size == (100, 300)
    corrected.close()


def test_path_is_not_kept():
    assert safe_filename("../../card.jpg") == "card.jpg"
    assert safe_filename("C:\\private\\card.jpg") == "card.jpg"


def test_export_opens_and_preserves_text_without_formulas():
    values = ["=1+1", "+44123456", "-1+1", "@SUM(A1)", "00123"]
    leads = [Lead(first_name=value, phone="00123456", email="a@example.com") for value in values]
    workbook = load_workbook(BytesIO(generate_excel(leads)))
    sheet = workbook["Leads"]
    assert tuple(cell.value for cell in sheet[1]) == HEADERS
    assert sheet.max_row == len(leads) + 1
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:G6"
    assert sheet["A1"].font.bold
    for index, value in enumerate(values, 2):
        assert sheet.cell(index, 1).value == value
        assert sheet.cell(index, 1).data_type == "s"
        assert sheet.cell(index, 6).value == "00123456"
        assert sheet.cell(index, 6).number_format == "@"
    workbook.close()
